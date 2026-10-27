"""Spreadsheets the user attaches in the chat (Excel .xlsx, CSV): parse, describe, query and
turn into inbox items (owner: track REPORT, file analysis).

Design (the same shape as the report agent, but over the user's file instead of the DB):
- An upload is parsed once into text cells and kept in memory for this user for a while
  (`UploadStore`). Nothing from the file is written to the database until the user confirms
  an import, and then only as inbox items they still approve one by one or in bulk.
- The agent reads a file through query tools, never by running code: `describe` (columns,
  types, examples), `rows` (a page), `query` (one SELECT over a table called `data`) and
  `prepare_import` (a preview of expense lines, written nowhere).
- `query` loads the sheet into a private in-memory SQLite database with an authorizer that
  only allows reading that one table, so a model-written query cannot touch anything else.
  Numbers there are floats: fine for analysis, and every money amount that becomes an
  expense goes through `statements.parse_amount` and exact minor units instead.
- File content is data, never instructions; the agent prompt says so and the chat API
  wraps the file list in a tagged block.

The store lives in process memory (one Cloud Run instance in the demo). Several instances
would need a shared store; files are small and short-lived, so that is a later change.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import secrets
import sqlite3
import threading
import time
from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import Decimal

import sqlglot
from sqlglot import exp

from app.services import statements
from app.services.errors import InvalidInput, NotFound
from app.services.money import MoneyError, format_minor, normalize_currency, to_minor

MAX_ROWS = 20_000
MAX_COLUMNS = 60
MAX_FILES_PER_USER = 5
MAX_DRAFTS_PER_USER = 5
TTL_SECONDS = 2 * 60 * 60
QUERY_ROW_LIMIT = 200
_HEADER_SCAN_ROWS = 20

SPREADSHEET_TYPES = {
    "text/csv",
    "application/csv",
    "text/plain",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[str | None]]  # cell text; numbers as plain decimals, dates as ISO
    types: list[str] = field(default_factory=list)  # "number" | "date" | "text" per column


@dataclass
class Upload:
    id: str
    user_id: str
    filename: str
    sheets: list[Sheet]
    created: float = field(default_factory=time.monotonic)

    def sheet(self, name: str = "") -> Sheet:
        if not name:
            return self.sheets[0]
        for s in self.sheets:
            if s.name.lower() == name.strip().lower():
                return s
        raise InvalidInput(f"no sheet {name!r}; sheets: {', '.join(s.name for s in self.sheets)}")


@dataclass
class ImportDraft:
    id: str
    user_id: str
    file_id: str
    filename: str
    lines: list[statements.StatementLine]
    raws: list[str]
    created: float = field(default_factory=time.monotonic)


class _Store:
    """Per-user, size-capped, expiring map of id -> object."""

    def __init__(self, per_user: int) -> None:
        self._items: OrderedDict[str, Upload | ImportDraft] = OrderedDict()
        self._per_user = per_user
        self._lock = threading.Lock()

    def put(self, item: Upload | ImportDraft) -> None:
        with self._lock:
            self._expire()
            self._items[item.id] = item
            mine = [k for k, v in self._items.items() if v.user_id == item.user_id]
            for old in mine[: max(0, len(mine) - self._per_user)]:
                del self._items[old]

    def get(self, user_id: str, item_id: str) -> Upload | ImportDraft:
        with self._lock:
            self._expire()
            item = self._items.get(item_id)
        # Another user's id looks exactly like a missing one.
        if item is None or item.user_id != user_id:
            raise NotFound("that file is no longer available; attach it again")
        return item

    def pop(self, user_id: str, item_id: str) -> Upload | ImportDraft:
        item = self.get(user_id, item_id)
        with self._lock:
            self._items.pop(item_id, None)
        return item

    def clear(self) -> None:
        with self._lock:
            self._items.clear()

    def _expire(self) -> None:
        now = time.monotonic()
        for key in [k for k, v in self._items.items() if now - v.created > TTL_SECONDS]:
            del self._items[key]


uploads = _Store(MAX_FILES_PER_USER)
drafts = _Store(MAX_DRAFTS_PER_USER)


# --- parsing --------------------------------------------------------------------------------


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        # repr gives the shortest exact form ("12.5", not "12.500000001").
        text = format(Decimal(repr(value)).normalize(), "f")
        return text
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _column_names(header: list[str | None]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(header):
        name = re.sub(r"\s+", " ", (raw or "").strip())[:60] or f"column_{index + 1}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name}_{seen[key]}"
        else:
            seen[key] = 1
        names.append(name)
    return names


def _table(name: str, raw_rows: list[list[str | None]]) -> Sheet | None:
    rows = [r[:MAX_COLUMNS] for r in raw_rows if any(c not in (None, "") for c in r)]
    if not rows:
        return None
    # The header is the first of the early rows with at least two filled cells (banks often
    # put a title and the account number above the table).
    header_at = next(
        (i for i, r in enumerate(rows[:_HEADER_SCAN_ROWS]) if sum(c is not None for c in r) >= 2), 0
    )
    header, body = rows[header_at], rows[header_at + 1 :]
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    header = header + [None] * (width - len(header))
    body = [r + [None] * (width - len(r)) for r in body[:MAX_ROWS]]
    sheet = Sheet(name=name, columns=_column_names(header), rows=body)
    sheet.types = [_infer_type([r[i] for r in body]) for i in range(width)]
    return sheet


def _infer_type(values: list[str | None]) -> str:
    filled = [v for v in values if v not in (None, "")][:500]
    if not filled:
        return "text"
    dates = sum(1 for v in filled if statements.parse_date(v) is not None)
    if dates >= 0.8 * len(filled):
        return "date"
    numbers = sum(1 for v in filled if statements.parse_amount(v) is not None)
    return "number" if numbers >= 0.8 * len(filled) else "text"


def _parse_xlsx(data: bytes) -> list[Sheet]:
    from openpyxl import load_workbook  # imported lazily: only chat uploads need it

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for broken files
        raise InvalidInput("this does not look like an Excel .xlsx file") from exc
    sheets = []
    try:
        for ws in book.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([_cell_text(v) for v in row])
                if len(rows) > MAX_ROWS + _HEADER_SCAN_ROWS:
                    break
            table = _table(ws.title, rows)
            if table:
                sheets.append(table)
    finally:
        book.close()
    return sheets


def _parse_csv(data: bytes) -> list[Sheet]:
    text = statements.decode_text(data)
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = [[c.strip() or None for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    table = _table("Sheet1", rows[: MAX_ROWS + _HEADER_SCAN_ROWS])
    return [table] if table else []


def parse_file(filename: str, data: bytes) -> list[Sheet]:
    """Excel (.xlsx) or CSV bytes -> non-empty sheets."""
    if not data:
        raise InvalidInput("the file is empty")
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        sheets = _parse_xlsx(data)
    elif name.endswith(".xls") and data[:4] == b"\xd0\xcf\x11\xe0":
        raise InvalidInput("old .xls files are not supported; save it as .xlsx or CSV")
    else:
        sheets = _parse_csv(data)
    if not sheets:
        raise InvalidInput("no table found in the file")
    return sheets


def store_upload(user_id: str, filename: str, data: bytes) -> Upload:
    upload = Upload(
        id=secrets.token_hex(8),
        user_id=user_id,
        filename=filename or "file",
        sheets=parse_file(filename, data),
    )
    uploads.put(upload)
    return upload


def get_upload(user_id: str, file_id: str) -> Upload:
    item = uploads.get(user_id, file_id)
    assert isinstance(item, Upload)
    return item


def summary_line(upload: Upload) -> str:
    """One line per sheet for the chat note: name, row count, columns."""
    parts = [f"{s.name} [{len(s.rows)} rows; columns: {', '.join(s.columns)}]" for s in upload.sheets]
    return f"{upload.filename} (file_id={upload.id}; sheets: {'; '.join(parts)})"


# --- describe / rows ------------------------------------------------------------------------


def describe(upload: Upload, sheet_name: str = "") -> dict:
    """Columns with type, fill count, examples and simple stats; plus every sheet's size."""
    sheet = upload.sheet(sheet_name)
    columns = []
    for i, name in enumerate(sheet.columns):
        values = [r[i] for r in sheet.rows if r[i] not in (None, "")]
        info: dict = {"name": name, "type": sheet.types[i], "filled": len(values), "examples": values[:3]}
        if sheet.types[i] == "number":
            numbers = [statements.parse_amount(v) for v in values]
            amounts = [n[0] for n in numbers if n is not None]
            if amounts:
                info.update(min=str(min(amounts)), max=str(max(amounts)), sum=str(sum(amounts)))
        elif sheet.types[i] == "date":
            days = sorted(d for v in values if (d := statements.parse_date(v)) is not None)
            if days:
                info.update(first=days[0].isoformat(), last=days[-1].isoformat())
        else:
            counts: dict[str, int] = {}
            for v in values:
                counts[v] = counts.get(v, 0) + 1
            info["distinct"] = len(counts)
            info["top"] = sorted(counts.items(), key=lambda kv: -kv[1])[:5]
        columns.append(info)
    return {
        "file": upload.filename,
        "sheet": sheet.name,
        "rows": len(sheet.rows),
        "columns": columns,
        "sheets": [{"name": s.name, "rows": len(s.rows)} for s in upload.sheets],
    }


def rows(upload: Upload, sheet_name: str = "", offset: int = 0, limit: int = 20) -> dict:
    sheet = upload.sheet(sheet_name)
    offset = max(0, offset)
    limit = max(1, min(limit, 50))
    page = sheet.rows[offset : offset + limit]
    return {
        "sheet": sheet.name,
        "columns": sheet.columns,
        "rows": [dict(zip(sheet.columns, r, strict=True)) for r in page],
        "offset": offset,
        "total_rows": len(sheet.rows),
        "next_offset": offset + len(page) if offset + len(page) < len(sheet.rows) else None,
    }


# --- query ----------------------------------------------------------------------------------

TABLE = "data"
_BLOCKED_FUNCTIONS = {"load_extension", "readfile", "writefile", "edit", "fts3_tokenizer"}


def _sql_value(value: str | None, kind: str) -> object:
    if value in (None, ""):
        return None
    if kind == "number":
        parsed = statements.parse_amount(value)
        return float(parsed[0]) if parsed else None
    if kind == "date":
        day = statements.parse_date(value)
        return day.isoformat() if day else value
    return value


def _validate(sql: str) -> str:
    try:
        statements_ = sqlglot.parse(sql, read="sqlite")
    except sqlglot.errors.ParseError as exc:
        raise InvalidInput(f"the SQL does not parse: {str(exc).splitlines()[0]}") from exc
    if len(statements_) != 1 or statements_[0] is None:
        raise InvalidInput("send exactly one SELECT statement")
    tree = statements_[0]
    if not isinstance(tree, (exp.Select, exp.Union)):
        raise InvalidInput("only SELECT is allowed")
    for table in tree.find_all(exp.Table):
        if table.name.lower() != TABLE:
            raise InvalidInput(f"the only table is {TABLE!r}")
    for func in tree.find_all(exp.Anonymous):
        if func.name.lower() in _BLOCKED_FUNCTIONS:
            raise InvalidInput(f"function {func.name} is not allowed")
    return sql


def _authorizer(
    action: int, arg1: str | None, arg2: str | None, _db: str | None, _trigger: str | None
) -> int:
    if action == sqlite3.SQLITE_SELECT:
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_READ and (arg1 or "").lower() == TABLE:
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_FUNCTION and (arg2 or "").lower() not in _BLOCKED_FUNCTIONS:
        return sqlite3.SQLITE_OK
    return sqlite3.SQLITE_DENY


def query(upload: Upload, sql: str, sheet_name: str = "", limit: int = QUERY_ROW_LIMIT) -> dict:
    """Run one SELECT over the sheet as table `data` (column names as in the header, quoted
    with double quotes when they have spaces). Dates are ISO text, so strftime('%Y-%m', x)
    works; numbers are REAL."""
    sheet = upload.sheet(sheet_name)
    _validate(sql)
    conn = sqlite3.connect(":memory:")
    try:
        cols = ", ".join(
            f'"{c.replace(chr(34), "")}" {"REAL" if t == "number" else "TEXT"}'
            for c, t in zip(sheet.columns, sheet.types, strict=True)
        )
        conn.execute(f"CREATE TABLE {TABLE} ({cols})")
        marks = ", ".join("?" for _ in sheet.columns)
        conn.executemany(
            f"INSERT INTO {TABLE} VALUES ({marks})",
            ([_sql_value(v, t) for v, t in zip(r, sheet.types, strict=True)] for r in sheet.rows),
        )
        conn.set_authorizer(_authorizer)
        try:
            cursor = conn.execute(sql)
            fetched = cursor.fetchmany(limit + 1)
        except sqlite3.Error as exc:
            raise InvalidInput(f"the query failed: {exc}") from exc
        columns = [d[0] for d in cursor.description or []]
    finally:
        conn.close()
    out_rows = [
        {c: (round(v, 2) if isinstance(v, float) else v) for c, v in zip(columns, r, strict=True)}
        for r in fetched[:limit]
    ]
    return {
        "columns": columns,
        "rows": out_rows,
        "row_count": len(out_rows),
        "truncated": len(fetched) > limit,
    }


# --- import ---------------------------------------------------------------------------------


def prepare_import(
    upload: Upload,
    *,
    date_column: str,
    description_column: str,
    amount_column: str,
    sheet_name: str = "",
    currency_column: str = "",
    currency: str = "TRY",
    spending_sign: str = "auto",
) -> tuple[ImportDraft, dict]:
    """Turn sheet rows into statement lines (spending only) without saving anything.

    `spending_sign`: "positive" when spending is written as positive amounts, "negative"
    when bank style (spending negative, income positive), "auto" picks the sign most rows
    have. Rows with the other sign (refunds, salary) and unreadable rows are skipped."""
    sheet = upload.sheet(sheet_name)
    index = {c.lower(): i for i, c in enumerate(sheet.columns)}

    def col(name: str, required: bool = True) -> int | None:
        if not name:
            if required:
                raise InvalidInput("a column name is missing")
            return None
        if name.strip().lower() not in index:
            raise InvalidInput(f"no column {name!r}; columns: {', '.join(sheet.columns)}")
        return index[name.strip().lower()]

    di, ti, ai, ci = (
        col(date_column),
        col(description_column),
        col(amount_column),
        col(currency_column, False),
    )
    try:
        default_currency = normalize_currency(currency or "TRY")
    except MoneyError as exc:
        raise InvalidInput(str(exc)) from exc

    parsed: list[tuple[int, dt.date | None, str, tuple[Decimal, str | None] | None, str | None]] = []
    for n, r in enumerate(sheet.rows):
        parsed.append(
            (
                n,
                statements.parse_date(r[di] or ""),
                (r[ti] or "").strip(),
                statements.parse_amount(r[ai] or ""),
                r[ci] if ci is not None else None,
            )
        )
    if spending_sign == "auto":
        negatives = sum(1 for p in parsed if p[3] and p[3][0] < 0)
        positives = sum(1 for p in parsed if p[3] and p[3][0] > 0)
        spending_sign = "negative" if negatives > positives else "positive"
    if spending_sign not in ("positive", "negative"):
        raise InvalidInput("spending_sign must be positive, negative or auto")

    lines: list[statements.StatementLine] = []
    raws: list[str] = []
    skipped: list[str] = []
    for n, day, text, amount, cell_currency in parsed:
        row_no = n + 1
        if day is None or amount is None or not text:
            skipped.append(f"row {row_no}: missing or unreadable date, description or amount")
            continue
        value, found_currency = amount
        if value == 0 or (value < 0) != (spending_sign == "negative"):
            skipped.append(f"row {row_no}: not spending (refund, income or zero)")
            continue
        try:
            line_currency = normalize_currency(cell_currency or found_currency or default_currency)
            minor = to_minor(str(abs(value)))
        except MoneyError:
            skipped.append(f"row {row_no}: unreadable amount or currency")
            continue
        lines.append(statements.StatementLine(day, text[:200], format_minor(minor), line_currency))
        raws.append(" | ".join(v or "" for v in sheet.rows[n]))

    if not lines:
        raise InvalidInput("no spending rows found with those columns")
    draft = ImportDraft(
        id=secrets.token_hex(8),
        user_id=upload.user_id,
        file_id=upload.id,
        filename=upload.filename,
        lines=lines,
        raws=raws,
    )
    drafts.put(draft)
    totals: dict[str, int] = {}
    for line in lines:
        totals[line.currency] = totals.get(line.currency, 0) + to_minor(line.amount)
    preview = {
        "status": "ready",
        "draft_id": draft.id,
        "file": upload.filename,
        "count": len(lines),
        "totals": [{"currency": c, "total": format_minor(m)} for c, m in sorted(totals.items())],
        "first": min(x.posted_on for x in lines).isoformat(),
        "last": max(x.posted_on for x in lines).isoformat(),
        "sample": [
            {
                "date": x.posted_on.isoformat(),
                "description": x.description,
                "amount": x.amount,
                "currency": x.currency,
            }
            for x in lines[:5]
        ],
        "skipped": len(skipped),
        "skipped_examples": skipped[:5],
    }
    return draft, preview


def get_draft(user_id: str, draft_id: str) -> ImportDraft:
    item = drafts.get(user_id, draft_id)
    assert isinstance(item, ImportDraft)
    return item
